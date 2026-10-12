"""Builds the assignment report by filling the course template (template.docx).

Run with a Python that has python-docx:   python knn_nb/report/build_report.py
Output: knn_nb/report/KNN_NaiveBayes_Report.docx
"""
from __future__ import annotations

from pathlib import Path

import docx
import pandas as pd
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).parent
KNN = HERE.parent
FIG = KNN / "figures"
OUT = HERE / "KNN_NaiveBayes_Report.docx"

d = docx.Document(HERE / "template.docx")
anchor = d.paragraphs[-1]                       # "< your assignment starts from here [good luck]>"
for r in anchor.runs:
    r.text = ""

fig_no = [0]
tab_no = [0]


# ---------------------------------------------------------------- helpers
def h1(text):
    d.add_paragraph(text, style="Heading 1")


def h2(text):
    d.add_paragraph(text, style="Heading 2")


def para(*parts, italic=False, size=None, align=None):
    """parts: str or (str, 'b') for bold."""
    p = d.add_paragraph()
    for part in parts:
        text, bold = (part, False) if isinstance(part, str) else (part[0], True)
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(6)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def point(lead, text):
    p = para((lead, "b"), " " + text)
    p.paragraph_format.left_indent = Inches(0.3)
    return p


def equation(text):
    p = para(text, align="center")
    for r in p.runs:
        r.font.name = "Cambria Math"
        r.italic = True
    return p


def figure(name, caption, explain):
    fig_no[0] += 1
    d.add_paragraph().add_run().add_picture(str(FIG / name), width=Inches(6.6))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = para((f"Figure {fig_no[0]}: ", "b"), caption, italic=True, size=9, align="center")
    cap.paragraph_format.space_after = Pt(4)
    para(("What the graph shows: ", "b"), explain)


def set_borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "999999")
        b.append(e)
    pr.append(b)


def shade(cell, hex_):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_)
    cell._tc.get_or_add_tcPr().append(sh)


def table(header, rows, caption, bold_rows=()):
    tab_no[0] += 1
    cap = para((f"Table {tab_no[0]}: ", "b"), caption, italic=True, size=9, align="center")
    cap.paragraph_format.space_after = Pt(2)
    t = d.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borders(t)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.text = ""
        run = c.paragraphs[0].add_run(str(h)); run.bold = True; run.font.size = Pt(8.5)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
        shade(c, "D9E2F3")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]
            c.text = ""
            run = c.paragraphs[0].add_run(str(v)); run.font.size = Pt(9)
            if i - 1 in bold_rows:
                run.bold = True
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
    d.add_paragraph()


def code(path: Path, start=None, end=None, title=None):
    lines = path.read_text(encoding="utf-8").splitlines()
    if start is not None:
        i0 = next(i for i, l in enumerate(lines) if l.startswith(start))
        i1 = next(i for i, l in enumerate(lines) if i > i0 and l.startswith(end)) if end else len(lines)
        lines = lines[i0:i1]
    if title:
        para((title, "b"))
    for text in ["#YOUR CODE STARTS HERE#", *lines, "#YOUR CODE ENDS HERE#"]:
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text if text.strip() else " ")        # spaces keep Python indentation exact
        r.font.name = "Courier New"; r.font.size = Pt(8.5)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        if text.startswith("#YOUR CODE"):
            r.bold = True; r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    d.add_paragraph()


pct = lambda v: f"{float(v) * 100:.1f}%"
pv = lambda v: "<0.001" if float(v) < 0.001 else f"{float(v):.3f}"
r1 = pd.read_csv(KNN / "results.csv", index_col=0)
r2 = pd.read_csv(KNN / "results_vol.csv", index_col=0)
r3 = pd.read_csv(KNN / "results_regime.csv", index_col=0)

# =========================================================================== 1
h1("1. Introduction and Real-Life Problem")
para("Investors, banks and fund managers make decisions every day based on what they think a stock will do next. "
     "Two questions matter most: (a) will the price go up or down, and (b) how much will it move (how risky it will be). "
     "In this assignment we built K-Nearest Neighbours (K-NN) and Gaussian Naive Bayes classifiers completely from scratch "
     "(using only NumPy, Pandas and Matplotlib) and used them to try to answer these questions honestly for five large "
     "technology stocks: Apple (AAPL), Microsoft (MSFT), NVIDIA (NVDA), Alphabet (GOOGL) and Meta (META), using about "
     "22 years of daily prices (January 2004 to October 2026).")
para(("Real-life problem solved: ", "b"),
     "predicting whether the next month will be a high-volatility (risky) or low-volatility (calm) month for a stock. "
     "This is a real risk-management problem: a fund that knows a stormy month is coming can hold less of that stock, "
     "and an options trader can price insurance correctly. Along the way we also tested the harder question of predicting "
     "the direction of tomorrow's price, and report honestly that it cannot be done with these tools.")

# =========================================================================== 2
h1("2. Technical Background")
h2("2.1 K-Nearest Neighbours (K-NN)")
para("K-NN is the simplest machine-learning idea: similar situations tend to have similar outcomes. "
     "It does not learn any equation. It simply stores every training example (this is why it is called a "
     "\"lazy learner\"). To classify a new example it does three steps:")
point("Step 1 - Distance:", "measure how far the new example is from every training example. We use Euclidean distance:")
equation("d(x, x′) = √( Σⱼ (xⱼ − x′ⱼ)² )")
point("Step 2 - Neighbours:", "keep the K training examples with the smallest distance.")
point("Step 3 - Vote / average:", "for classification, the most common class among the K neighbours is the prediction. "
      "For a number (regression), the prediction is the average of the neighbours' values. We use the average because it "
      "also gives a range: the spread of the neighbours' outcomes, which we plot as a shaded band.")
para(("Important details. ", "b"),
     "(1) Scaling: distance is dominated by features with large numbers, so every feature is standardised "
     "(z = (x − mean) / standard deviation) using the mean and standard deviation of the training data only. "
     "(2) Choosing K: a small K follows noise (over-fitting); a large K smooths too much (under-fitting). "
     "We choose K on a separate validation period, never on the test data.")

h2("2.2 Gaussian Naive Bayes")
para("Naive Bayes is based on Bayes' theorem. It asks: given the features I see today, what is the probability of each class?")
equation("P(class | x₁ … xₙ)  ∝  P(class) × P(x₁ | class) × P(x₂ | class) × … × P(xₙ | class)")
point("Prior P(class):", "how common each class is in the training data (for example, the share of high-volatility months).")
point("Likelihood P(xⱼ | class):", "how typical a feature value is inside each class. For continuous features we assume a "
      "normal (Gaussian) distribution, so for every class and feature we only need to learn a mean μ and a variance σ²:")
equation("P(xⱼ | c) = 1/√(2πσ²꜀ⱼ) · exp( −(xⱼ − μ꜀ⱼ)² / (2σ²꜀ⱼ) )")
point("\"Naive\" assumption:", "features are treated as independent of each other inside each class, which is what allows "
      "the likelihoods to be simply multiplied. This is rarely exactly true, and we will see its effect in the results.")
para("Two implementation details: we add log-probabilities instead of multiplying (to avoid numbers becoming too small "
     "for the computer, called underflow), and we add a tiny variance floor so we never divide by zero.")

h2("2.3 K-NN vs Naive Bayes")
table(["", "K-NN", "Naive Bayes"],
      [["What it learns", "Nothing - stores the data", "Class priors, means, variances"],
       ["Prediction speed", "Slow (compares to all data)", "Very fast"],
       ["Main assumption", "Similar inputs → similar outputs", "Features independent given the class"],
       ["Needs feature scaling", "Yes", "No"],
       ["Output", "Class vote / average and spread", "A probability for each class"]],
      "Comparison of the two classifiers")

# =========================================================================== 3
h1("3. Dataset and Pre-processing")
para(("Source. ", "b"),
     "Daily adjusted closing prices were downloaded from Yahoo Finance (with the yfinance library) for AAPL, MSFT, NVDA, "
     "GOOGL and META from 2 January 2004 to 9 October 2026 (5,729 trading days). Adjusted prices already account for "
     "stock splits and dividends, so a price jump caused by a split does not look like a real return. "
     "GOOGL starts trading in August 2004 and META in May 2012, so they have shorter histories.")
para(("Cleaning. ", "b"),
     "All stocks were aligned to the same New York Stock Exchange trading calendar. Gaps of at most 2 days were "
     "forward-filled (the last known price is carried forward); nothing was ever filled backwards, because that would "
     "copy future prices into the past. Days before a stock was listed were left empty.")
para(("Returns. ", "b"), "Daily return r(t) = P(t) / P(t−1) − 1. Volatility = standard deviation of daily returns, "
     "multiplied by √252 so it is expressed per year (annualised %), which is how the industry quotes it.")
para(("Features (inputs) - all known at the close of day t: ", "b"))
table(["Feature", "Meaning", "Used in"],
      [["ret1, ret5, ret21, ret63", "Return over the last 1, 5, 21 and 63 trading days", "Problem 1"],
       ["abs1", "Size of today's move |r(t)|", "Problem 2"],
       ["vol5, vol21, vol63, vol252", "Volatility over the last 5, 21, 63 and 252 days", "Problems 2 & 3"],
       ["ret21", "Return over the last month (falling prices tend to raise risk)", "Problems 2 & 3"]],
      "Input features")
para(("Targets (what we predict). ", "b"),
     "Problem 1: will tomorrow be an up day (r(t+1) > 0)? Problem 2: will tomorrow's move |r(t+1)| be bigger than "
     "normal? Problem 3: will the volatility of the next 21 trading days be above this stock's normal (median) level?")
para(("Preventing leakage. ", "b"),
     "The biggest mistake in financial machine learning is letting information from the future leak into training. "
     "We prevented it in four ways: (1) data was split by time, never shuffled; (2) scaling statistics and the \"normal\" "
     "threshold were computed from training data only; (3) a training example was only used once its outcome had fully "
     "happened before the prediction date (for 21-day targets this creates a 21-day gap, called an embargo); "
     "(4) K was selected on an earlier validation period, and the test period was evaluated only once.")

# =========================================================================== 4
h1("4. What We Were Trying to Find")
para("We asked three questions, from hardest to easiest:")
point("Problem 1 - Direction:", "Can K-NN or Naive Bayes predict if a tech stock will go up or down tomorrow?")
point("Problem 2 - Size of tomorrow's move:", "Can they predict if tomorrow will be a big or a small move?")
point("Problem 3 - Next month's risk regime:", "Can they predict if the next month will be high- or low-volatility?")
para("For each question the honest test is not \"is accuracy above 50%?\" but \"does the model beat a simple rule "
     "that needs no machine learning?\" We used two such baselines:")
point("Always-majority:", "always guess the most common class of the test period (chosen in hindsight, so it is a generous baseline).")
point("Persistence:", "assume the future looks like the present (for example, \"next month will be as volatile as this month\").")
para("To check whether a difference is real or just luck we used a one-sided z-test on the accuracy: "
     "z = (accuracy − baseline) / √(baseline × (1 − baseline) / n). A p-value below 0.05 means the improvement is "
     "unlikely to be luck.")

# =========================================================================== 5
h1("5. How We Did It (Method)")
para(("Walk-forward testing. ", "b"),
     "Instead of training once, the models are re-trained at the start of every month on all data available up to that "
     "day, then used to predict. Then we step forward one month and repeat. This is exactly how a model would be used in "
     "real life, and it means the model never sees the period it is predicting.")
table(["", "Problem 1: Direction", "Problem 2: Next-day size", "Problem 3: Next-month regime"],
      [["Target", "Up / down tomorrow", "|r(t+1)| above median", "21-day vol above median"],
       ["Training data", "2004 → day before", "2004 → day before", "2004 → finished 21-day outcomes"],
       ["K chosen on", "2025", "2025", "2008-2011"],
       ["Chosen K", "200", "400", "100"],
       ["Test period", "Jan-Oct 2026", "Jan-Oct 2026", "Jan 2012-Sep 2026"],
       ["Predictions", "193 days × 5 stocks", "193 days × 5 stocks", "843 stock-months"]],
      "Set-up of the three experiments")
para("K was chosen from {5, 10, 20, 50, 100, 200, 400} by the lowest mean squared error of K-NN's prediction on the "
     "validation period. In Problem 3 one prediction is made per stock per month (on the first trading day), so the "
     "21-day outcomes do not overlap and each prediction is a separate, independent test.")
para("Problems 1 and 2 were tested only on 2026. After seeing that they failed, we checked why (Section 7) and only then "
     "designed Problem 3, fixing all its settings before running it. The longer 2012-2026 test was chosen because it "
     "contains several calm and stormy periods, which a one-year test does not.")

# =========================================================================== 6
h1("6. Results")
h2("6.1 Problem 1 - Predicting tomorrow's direction (Jan-Oct 2026)")
table(["Stock", "K-NN acc.", "Naive Bayes acc.", "Always-up acc.", "Buy & hold", "Follow K-NN", "Follow NB"],
      [[t, pct(r1.loc[t, "KNN acc"]), pct(r1.loc[t, "NB acc"]), pct(r1.loc[t, "best-constant acc"]),
        f"{r1.loc[t, 'hold %']:+.1f}%", f"{r1.loc[t, 'KNN signal %']:+.1f}%", f"{r1.loc[t, 'NB signal %']:+.1f}%"]
       for t in r1.index],
      "Direction accuracy and 2026 returns of following each model (no trading costs)")
para("Neither model beat the simple baseline on any stock (all p-values ≥ 0.50). K-NN was below 50% on every stock. "
     "Following the models' signals (buy when they say up, short-sell when they say down) lost money compared to just "
     "holding the stock.")
figure("g1_knn_shadow_vs_actual.png", "K-NN predicted next-day return (blue line and shaded ±1 standard deviation band) vs actual return (black), 2026",
       "The blue predicted line stays almost flat near zero, while the actual returns (black) jump between roughly −9% "
       "and +15%. The model has learned that on average tomorrow is close to zero, and that is all it knows. The shaded "
       "band is simply the normal daily noise of the stock, so the actual line falls inside it about 70% of the time by "
       "chance, not because the model is skilful. K-NN's error (MAE) was slightly worse than just predicting 0% every day.")
figure("g2_nb_prob_vs_actual.png", "Naive Bayes probability of an up day (orange) vs actual up (1) / down (0) days, 2026",
       "The orange area shows how far Naive Bayes' probability of an up day is above or below 50%. It stays very close "
       "to 50% almost every day, which means the model is unsure. The black dots (actual up or down) are scattered "
       "randomly with no relation to the orange area: high-probability days are not more often up.")
figure("g3_follow_vs_hold.png", "Growth of $1 in 2026: holding the stock (black) vs following each model's signal",
       "If the models had real skill, the coloured lines would end above the black line. Instead, K-NN's line ends "
       "below on all five stocks (as low as −61% for META) and Naive Bayes is mixed. This is the money view of the same "
       "result: no usable prediction of direction.")

h2("6.2 Problem 2 - Predicting the size of tomorrow's move (Jan-Oct 2026)")
table(["Stock", "K-NN acc.", "Naive Bayes acc.", "Always-majority", "Persistence"],
      [[t, pct(r2.loc[t, "KNN acc"]), pct(r2.loc[t, "NB acc"]), pct(r2.loc[t, "majority acc"]),
        pct(r2.loc[t, "persistence acc"])] for t in r2.index]
      + [["All pooled", "54.9%", "52.8%", "54.6%", "51.8%"]],
      "Accuracy of calling a big vs small move tomorrow", bold_rows=(5,))
figure("g6_vol_accuracy_vs_baselines.png", "Big-vs-small day accuracy in 2026: models (blue, orange) vs baselines (grey)",
       "For every stock the coloured bars (models) are about the same height as, or lower than, the grey bars "
       "(baselines). Pooled over all stocks K-NN reached 54.9% against 54.6% for always guessing the majority - "
       "no real difference. Predicting the size of one single day was also not possible in 2026.")

h2("6.3 Problem 3 - Predicting next month's volatility regime (2012-2026)")
rows3 = []
for t in r3.index:
    rows3.append([t, int(r3.loc[t, "n"]), pct(r3.loc[t, "majority"]), pct(r3.loc[t, "persistence"]),
                  pct(r3.loc[t, "K-NN"]), pct(r3.loc[t, "Naive Bayes"]),
                  pv(r3.loc[t, "p K-NN>maj"]), pv(r3.loc[t, "p NB>maj"])])
table(["Stock", "Months", "Majority", "Persist.", "K-NN", "NB", "p K-NN", "p NB"],
      rows3, "High vs low next-month volatility: accuracy, and p-value that the model beats always-majority "
             "(NB = Naive Bayes)", bold_rows=(5,))
table(["Measure (all stocks pooled)", "Value"],
      [["K-NN error on the volatility level (MAE)", f"{r3.loc['ALL (pooled)', 'K-NN MAE (vol pts)']:.1f} vol points"],
       ["Persistence error on the volatility level (MAE)", f"{r3.loc['ALL (pooled)', 'persistence MAE (vol pts)']:.1f} vol points"],
       ["p-value: K-NN accuracy > persistence", f"{r3.loc['ALL (pooled)', 'p K-NN>persist']:.2f}"],
       ["p-value: Naive Bayes accuracy > persistence", f"{r3.loc['ALL (pooled)', 'p NB>persist']:.2f}"],
       ["Actual inside K-NN's 25-75% band (expected 50%)", pct(r3.loc["ALL (pooled)", "band coverage"])]],
      "Extra measures for Problem 3")
para("Here both models clearly work: 66% (K-NN) and 68% (Naive Bayes) against 55% for always-majority, with p-values "
     "below 0.001. Four of the five stocks are significant on their own (AAPL is the exception, explained in Section 7). "
     "However, the simple persistence rule reaches 66% too, so the models do not clearly beat it on accuracy "
     "(p = 0.33 and 0.06). K-NN's real extra value is in the size of the forecast: its volatility estimate is about 17% "
     "more accurate than persistence (9.2 vs 11.1 points). In 2026 alone (45 predictions) no method was significant.")
figure("g7_regime_process.png", "How every prediction in Problem 3 is made",
       "The blue block is the training data, the red block is the 21-day gap (embargo) and the green block is the next "
       "month we are predicting. Only past days whose 21-day result had completely finished before the prediction date "
       "are used to train, so the model cannot peek at the future. After each prediction we move forward one month, "
       "re-train and repeat.")
figure("g8_regime_knn_shadow_vs_actual.png", "K-NN predicted next-month volatility (blue line and shaded range) vs actual (black), 2012-2026",
       "Each step is one month. The blue line and shaded band are K-NN's forecast; the black line is what really "
       "happened; the red dashed line is the stock's normal level, which separates high and low months. The blue line "
       "follows the big waves: it rises in 2020 (COVID crash), stays high through 2022 (interest-rate shock for tech "
       "stocks) and falls in calm years such as 2017. What it cannot do is predict the first month of a sudden shock - "
       "the tall black spikes (e.g. March 2020, NVIDIA's November 2016 earnings jump) appear before the blue line "
       "reacts. The actual value falls inside the 25-75% band 43% of the time, slightly less than the 50% you would "
       "expect, because sudden shocks are bigger than anything the neighbours suggest.")
figure("g9_regime_nb_prob_vs_actual.png", "Naive Bayes probability of a high-volatility month (orange) vs actual high (red) / low (green) months",
       "Where the orange area is above the dotted 50% line the model predicts a HIGH month. The long orange blocks "
       "(2020, 2022, 2025) sit above long runs of red marks (actually high months), and the flat low periods "
       "(2013-2014, 2017) sit above green marks (actually low months), so the model is right most of the time. "
       "The orange area is almost always close to 0 or 1: Naive Bayes is very sure of itself, even when it is wrong.")
figure("g10_regime_scorecard.png", "Scorecard: accuracy of each method vs the two baselines",
       "Every coloured bar (K-NN, Naive Bayes) is well above the light-grey bar (always-majority), which shows there is "
       "something real to predict. But the dark-grey bar (persistence) is almost as tall as the model bars, which shows "
       "that most of the skill comes from the simple fact that volatility persists. The right-most group (2026 only) "
       "shows that in one calm year the advantage disappears.")
figure("g11_regime_trust_check.png", "Trust check: do the predictions line up with what really happened?",
       "Left: all 843 predictions sorted into 5 equal groups from \"calmest\" to \"stormiest\" according to K-NN. "
       "The months K-NN called calmest really were calm (0.80× the normal level) and the months it called stormiest "
       "really were stormy (1.51× normal) - a clean staircase means the ranking is trustworthy. "
       "Right: when Naive Bayes says 0-20% chance of a high month, 31% were actually high; when it says 80-100%, 75% "
       "were. The direction is right, but the probabilities are too extreme (over-confident) compared to the dotted line.")

# =========================================================================== 7
h1("7. Comparison with Real Life - Why These Percentages?")
h2("7.1 Why direction is about 50% (a coin flip)")
para("Stock prices are watched by millions of traders and computer programs. If yesterday's pattern really predicted "
     "tomorrow's direction, these traders would buy or sell on it immediately, and the pattern would disappear. "
     "This is the idea of an efficient market. What moves a price tomorrow is new information - an earnings report, a "
     "product launch, an interest-rate decision - which by definition is not in yesterday's data. So a model that only "
     "sees past prices cannot do better than a coin flip, and that is exactly what we found.")
para("The \"always up\" baseline is a bit above 50% (51-53%) because tech stocks rose over time: on average there are "
     "slightly more up days than down days. A model that sometimes says \"down\" loses this small advantage, which is why "
     "following the signals in Figure 3 did worse than just holding. K-NN's accuracy below 50% is mostly bad luck in a "
     "small sample: with 193 days, normal random variation is about ±3.6 percentage points.")

h2("7.2 Why the size of one single day is also unpredictable")
para("Over the whole history, a calm period is followed by calm days and a stormy period by stormy days. But one single "
     "day is dominated by random news, so the link from today to tomorrow is weak (correlation about 0.25). In 2026 it was "
     "even weaker (about 0): 2026 was one steady market with no big crisis, and the volatility of the five stocks moved "
     "only 2-4 times between its lowest and highest month, compared with 14-21 times over 2005-2025. If the thing you "
     "predict hardly changes, there is nothing to forecast.")

h2("7.3 Why next month's risk is predictable (66-68%)")
para("This is a well-known real-world fact called volatility clustering: periods of uncertainty last for weeks or months. "
     "When COVID hit in March 2020, markets stayed wild for months; during 2022, rising interest rates kept tech stocks "
     "nervous for the whole year. Averaging over a month removes the day-to-day noise, so the link between this month's "
     "and next month's volatility is strong (correlation about 0.5-0.6). The whole finance industry relies on this fact: "
     "the VIX \"fear index\", option prices, GARCH models and bank risk limits all assume that today's risk tells you "
     "about tomorrow's risk.")
para(("Why not higher than about 70%? ", "b"),
     "(1) Shocks arrive without warning - the first month of every crisis is missed, as the black spikes in Figure 6 show. "
     "(2) Many months are close to the normal level, and for these a small error flips the high/low label, so they are "
     "almost a coin flip.")
para(("Why the models only tie with persistence? ", "b"),
     "All the useful information is \"how volatile has it been recently\", and the persistence rule already uses exactly "
     "that. K-NN and Naive Bayes do not have any new information; they can only use it a little more smoothly. That is "
     "why K-NN's forecast of the volatility level is better (it averages 100 similar past months instead of copying one "
     "noisy month), while the high/low accuracy is about the same.")
para(("Why is Naive Bayes over-confident? ", "b"),
     "Its \"naive\" assumption says the five features are independent. In reality volatility over 5, 21, 63 and 252 days "
     "are strongly related - they all describe the same thing. Naive Bayes therefore counts the same evidence several "
     "times, which pushes its probabilities towards 0% or 100% (Figures 7 and 9). The class decision is still good, but "
     "the probability should not be trusted as a real percentage.")
para(("Why does AAPL have a 68% majority baseline? ", "b"),
     "Apple was a much riskier stock in 2004-2011 (when it was a smaller company) than after 2012. The \"normal\" level is "
     "learned from all the past, so most months after 2012 count as \"low\", and simply guessing \"low\" is already right "
     "68% of the time. That leaves little room for a model to improve, which is why AAPL is the only stock where the "
     "improvement is not significant.")

# =========================================================================== 8
h1("8. Conclusion")
para("We implemented K-NN and Gaussian Naive Bayes from scratch and tested them honestly with walk-forward testing, a "
     "leakage gap, simple baselines and significance tests on 22 years of tech-stock data. The findings are:")
point("1.", "Tomorrow's direction cannot be predicted from past prices (≈50%, no model beat the baseline).")
point("2.", "The size of tomorrow's move could not be predicted in 2026 either.")
point("3.", "Next month's risk regime can be predicted: 66-68% accuracy vs 55% for guessing, highly significant over "
      "2012-2026. Most of this comes from volatility persisting; K-NN adds a more accurate size forecast, and Naive Bayes "
      "gives good decisions but over-confident probabilities.")
para("The main lesson is that machine learning cannot find a signal that is not in the data. The value of these "
     "models in finance is in measuring and forecasting risk, not in predicting whether a price will go up or down.")

# =========================================================================== 9
h1("9. K-NN Worked on Paper")
para("[Insert the scanned pages of the hand-worked K-NN example here: a dataset with 2 features and at least 20 "
     "examples, showing the distance calculation, the K nearest neighbours and the final vote.]", italic=True)

# =========================================================================== 10
h1("10. Code")
para("All code uses only NumPy, Pandas and Matplotlib. The full files (knn.py, naive_bayes.py, run.py, run_vol.py, "
     "run_regime.py) and the dataset are attached with the submission. The core parts are shown below.")
code(KNN / "knn.py", title="10.1 knn.py - K-NN from scratch")
code(KNN / "naive_bayes.py", title="10.2 naive_bayes.py - Gaussian Naive Bayes from scratch")
code(KNN / "run_regime.py", start="def build", end="def pval_above", title="10.3 Pre-processing: features and target (run_regime.py)")
code(KNN / "run_regime.py", start="def walk", end="def score", title="10.4 Walk-forward training and prediction (run_regime.py)")

d.save(OUT)
print("saved", OUT)
